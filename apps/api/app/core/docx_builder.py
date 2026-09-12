"""DOCX documents for filtered on-demand admin exports."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from app.services.export_content import ExportFilters, FilteredExportDocument


def _filters_lines(filters: ExportFilters) -> list[str]:
    lines = [
        f"Status: {filters.status or 'All'}",
        f"Keyword: {filters.keyword or 'None'}",
        f"Created from: {filters.created_from or 'Any'}",
        f"Created to: {filters.created_to or 'Any'}",
    ]
    return lines


def build_filtered_export_docx(document: FilteredExportDocument) -> bytes:
    doc = Document()
    doc.add_heading("Filtered Report Export", level=1)
    doc.add_paragraph(f"Exported at: {document.exported_at}")
    doc.add_paragraph(f"Reports included: {document.report_count}")
    if document.truncated:
        doc.add_paragraph(
            f"Export limited to {document.export_limit} reports. "
            "Narrow filters for a complete set."
        )
    doc.add_heading("Applied filters", level=2)
    for line in _filters_lines(document.filters):
        doc.add_paragraph(line)

    if not document.reports:
        doc.add_heading("Results", level=2)
        doc.add_paragraph("No reports matched these filters.")
    else:
        for index, report in enumerate(document.reports, start=1):
            doc.add_heading(f"Report {index}: {report.id}", level=2)
            doc.add_paragraph(f"Type: {report.report_type}")
            doc.add_paragraph(f"Status: {report.status}")
            doc.add_paragraph(f"Severity: {report.severity}")
            doc.add_paragraph(f"Created: {report.created_at}")
            doc.add_paragraph(f"Updated: {report.updated_at}")
            doc.add_paragraph(f"Incident date: {report.incident_date or 'N/A'}")
            doc.add_paragraph(f"Incident location: {report.incident_location or 'N/A'}")
            if report.categories:
                doc.add_paragraph(f"Categories: {', '.join(report.categories)}")
            if report.reported_member_name:
                doc.add_paragraph(f"Named member: {report.reported_member_name}")
            doc.add_paragraph("Description:")
            doc.add_paragraph(report.description)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
